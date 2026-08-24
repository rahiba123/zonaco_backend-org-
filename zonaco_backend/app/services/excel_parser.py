"""Parser for Excel FAQ Knowledge Base with link extraction and document normalization."""

import hashlib
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional
import openpyxl
from app.config import get_settings
from app.utils.logger import logger

# Regex for extracting HTTP(S) URLs and www links
URL_REGEX = re.compile(
    r'(?:https?://[^\s<>"\')]+|www\.[^\s<>"\')]+)',
    re.IGNORECASE
)


@dataclass
class FAQEntry:
    """Structured representation of a single row in the FAQ knowledge base."""
    doc_id: str
    category: str
    s_no: str
    topic_type: str
    question: str
    response: str
    faq_intent: str
    details: str
    links: List[str] = field(default_factory=list)

    @property
    def formatted_text(self) -> str:
        """Formatted QA text optimized for vector embedding and retrieval."""
        return f"Question: {self.question}\nAnswer: {self.response}"

    def to_metadata(self) -> Dict[str, str]:
        """Convert entry to ChromaDB-compatible metadata dictionary (string values)."""
        import json
        return {
            "doc_id": self.doc_id,
            "category": self.category,
            "s_no": str(self.s_no),
            "topic_type": self.topic_type,
            "faq_intent": self.faq_intent,
            "question": self.question,
            "response": self.response,
            "details": self.details,
            "links": json.dumps(self.links),
        }


class ExcelFAQParser:
    """Extracts and normalizes FAQ knowledge base entries from Excel workbook."""

    KNOWN_CATEGORIES = [
        "CARD SERVICES",
        "EFTS",
        "CREDIT",
        "BILL MUSTER",
        "INTERNET BANKING",
        "GENERAL BANKING",
        "MOBILE BANKING",
        "ZANACO XPRESS",
    ]

    def __init__(self, file_path: Optional[str] = None):
        settings = get_settings()
        self.file_path = file_path or self._resolve_excel_path(settings.FAQ_EXCEL_PATH)

    @staticmethod
    def _resolve_excel_path(configured_path: str) -> str:
        """Locate the Excel file looking in root directory or configured path."""
        candidates = [
            configured_path,
            "Chat Bot Intents with Links (1).xlsx",
            "Chat_Bot_Intents_with_Links.xlsx",
            os.path.join(os.path.dirname(__file__), "..", "..", configured_path),
            os.path.join(os.path.dirname(__file__), "..", "..", "Chat Bot Intents with Links (1).xlsx"),
            os.path.join(os.path.dirname(__file__), "..", "..", "Chat_Bot_Intents_with_Links.xlsx"),
        ]
        for candidate in candidates:
            if candidate and os.path.exists(candidate):
                return os.path.abspath(candidate)
        return configured_path

    @staticmethod
    def extract_urls(text: Optional[str]) -> List[str]:
        """Extract all valid URLs from a text block, standardizing http scheme."""
        if not text or not isinstance(text, str):
            return []
        matches = URL_REGEX.findall(text)
        urls = []
        for match in matches:
            cleaned = match.rstrip(".,;:)")
            if cleaned.startswith("www."):
                cleaned = "https://" + cleaned
            if cleaned not in urls:
                urls.append(cleaned)
        return urls

    @staticmethod
    def generate_doc_id(category: str, s_no: str, question: str) -> str:
        """Generate deterministic SHA256 ID to ensure idempotent ingestion."""
        raw_key = f"{category.strip().upper()}::{str(s_no).strip()}::{question.strip().lower()}"
        return hashlib.sha256(raw_key.encode("utf-8")).hexdigest()

    def parse_all(self) -> List[FAQEntry]:
        """Parse all sheets in the Excel workbook and return normalized entries."""
        if not os.path.exists(self.file_path):
            logger.error(f"FAQ Excel file not found at: {self.file_path}")
            raise FileNotFoundError(f"FAQ Excel file not found: {self.file_path}")

        logger.info(f"Opening FAQ workbook at: {self.file_path}")
        wb = openpyxl.load_workbook(self.file_path, data_only=True)
        entries: List[FAQEntry] = []

        for sheet_name in wb.sheetnames:
            category_clean = sheet_name.strip().upper()
            sheet = wb[sheet_name]
            sheet_entries = self._parse_sheet(sheet, category_clean)
            logger.info(f"Parsed {len(sheet_entries)} FAQ entries from sheet '{sheet_name}'")
            entries.extend(sheet_entries)

        wb.close()
        logger.info(f"Total FAQ entries parsed across all categories: {len(entries)}")
        return entries

    def _parse_sheet(self, sheet, category: str) -> List[FAQEntry]:
        """Parse a single Excel worksheet into FAQEntry objects."""
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        # Find header row
        header_row_idx = None
        col_indices: Dict[str, int] = {}

        for r_idx, row in enumerate(rows[:10]):
            row_str = [str(c).strip().lower() if c is not None else "" for c in row]
            if any("question" in c for c in row_str) or any("s. no" in c for c in row_str):
                header_row_idx = r_idx
                for c_idx, cell_value in enumerate(row_str):
                    if "s. no" in cell_value or "s.no" in cell_value or "sno" in cell_value:
                        col_indices["s_no"] = c_idx
                    elif "type" in cell_value:
                        col_indices["topic_type"] = c_idx
                    elif "question" in cell_value:
                        col_indices["question"] = c_idx
                    elif "response" in cell_value:
                        col_indices["response"] = c_idx
                    elif "intent" in cell_value:
                        col_indices["faq_intent"] = c_idx
                    elif "details" in cell_value or "detail" in cell_value:
                        col_indices["details"] = c_idx
                break

        if header_row_idx is None:
            logger.warning(f"Could not identify header row in sheet '{category}'")
            return []

        sheet_entries: List[FAQEntry] = []
        for row in rows[header_row_idx + 1:]:
            # Check if row is empty or spacer
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            def get_val(key: str, default: str = "") -> str:
                idx = col_indices.get(key)
                if idx is not None and idx < len(row) and row[idx] is not None:
                    return str(row[idx]).strip()
                return default

            question = get_val("question")
            response = get_val("response")
            s_no = get_val("s_no", str(len(sheet_entries) + 1))
            topic_type = get_val("topic_type")
            faq_intent = get_val("faq_intent")
            details = get_val("details")

            # Must have at least question or response
            if not question and not response:
                continue

            links = self.extract_urls(details)
            doc_id = self.generate_doc_id(category, s_no, question)

            entry = FAQEntry(
                doc_id=doc_id,
                category=category,
                s_no=s_no,
                topic_type=topic_type,
                question=question,
                response=response,
                faq_intent=faq_intent,
                details=details,
                links=links
            )
            sheet_entries.append(entry)

        return sheet_entries

    def get_categories_overview(self) -> List[Dict[str, any]]:
        """Get overview list of categories and their question counts."""
        entries = self.parse_all()
        counts: Dict[str, int] = {}
        for entry in entries:
            counts[entry.category] = counts.get(entry.category, 0) + 1

        overview = []
        # Maintain standard category order
        for cat in self.KNOWN_CATEGORIES:
            if cat in counts:
                overview.append({
                    "name": cat,
                    "question_count": counts[cat],
                    "description": f"Frequently asked questions and guides for {cat.title()}"
                })
        for cat, cnt in counts.items():
            if cat not in self.KNOWN_CATEGORIES:
                overview.append({
                    "name": cat,
                    "question_count": cnt,
                    "description": f"Frequently asked questions for {cat.title()}"
                })
        return overview

    def get_questions_by_category(self, category: str) -> List[FAQEntry]:
        """Fetch all predefined questions under a specific category."""
        target_cat = category.strip().upper()
        entries = self.parse_all()
        return [e for e in entries if e.category.upper() == target_cat]
