"""Parser for user-uploaded documents (PDF / DOCX / TXT / XLSX / XLS / CSV / MD) with chunking for embedding.

Mirrors the structure of `excel_parser.py` so it fits the existing service layer:
each parsed chunk becomes a small dataclass with a deterministic doc_id and a
`to_metadata()` method that ChromaDB can store alongside the embedding.
"""

import hashlib
import os
from dataclasses import dataclass, field
from typing import Dict, List

from app.utils.exceptions import ChatbotException
from app.utils.logger import logger

# Chunking parameters — tune based on your embedding model's context window.
CHUNK_SIZE = 800       # characters per chunk
CHUNK_OVERLAP = 150    # overlap between consecutive chunks
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx", ".xls", ".csv", ".md", ".markdown"}
MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB safety cap


class UnsupportedFileTypeException(ChatbotException):
    """Raised when an uploaded file's extension isn't supported."""

    def __init__(self, filename: str):
        super().__init__(
            message=f"Unsupported file type for '{filename}'. Allowed: {sorted(ALLOWED_EXTENSIONS)}",
            status_code=400,
            details={"filename": filename}
        )


class DocumentParsingException(ChatbotException):
    """Raised when text extraction from an uploaded document fails."""

    def __init__(self, filename: str, reason: str):
        super().__init__(
            message=f"Failed to parse document '{filename}': {reason}",
            status_code=422,
            details={"filename": filename, "reason": reason}
        )


@dataclass
class DocumentChunk:
    """A single chunk of a user-uploaded document, ready for embedding."""
    doc_id: str
    session_id: str
    source_filename: str
    chunk_index: int
    text: str

    def to_metadata(self) -> Dict[str, str]:
        """Convert chunk to ChromaDB-compatible metadata dictionary."""
        return {
            "doc_id": self.doc_id,
            "session_id": self.session_id,
            "source_filename": self.source_filename,
            "chunk_index": str(self.chunk_index),
            "source": "user_upload",
        }


class DocumentParser:
    """Extracts and chunks text from uploaded PDF/DOCX/TXT/XLSX/XLS/CSV/MD files."""

    def validate(self, filename: str, file_bytes: bytes) -> None:
        """Raise if the file type or size isn't acceptable. Call before parsing."""
        ext = os.path.splitext(filename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise UnsupportedFileTypeException(filename)
        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            raise DocumentParsingException(
                filename, f"File exceeds {MAX_FILE_SIZE_BYTES // (1024*1024)}MB limit."
            )

    def extract_text(self, filename: str, file_bytes: bytes) -> str:
        """Dispatch to the right extractor based on file extension."""
        ext = os.path.splitext(filename)[1].lower()
        try:
            if ext == ".pdf":
                return self._extract_pdf(file_bytes)
            if ext == ".docx":
                return self._extract_docx(file_bytes)
            if ext in (".txt", ".md", ".markdown"):
                return file_bytes.decode("utf-8", errors="ignore")
            if ext in (".xlsx", ".xls"):
                return self._extract_excel(file_bytes, filename)
            if ext == ".csv":
                return self._extract_csv(file_bytes)
            raise UnsupportedFileTypeException(filename)
        except UnsupportedFileTypeException:
            raise
        except Exception as exc:
            logger.exception(f"Text extraction failed for '{filename}': {exc}")
            raise DocumentParsingException(filename, str(exc))

    def _extract_pdf(self, file_bytes: bytes) -> str:
        from io import BytesIO
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)

    def _extract_docx(self, file_bytes: bytes) -> str:
        from io import BytesIO
        import docx

        doc = docx.Document(BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    def _extract_excel(self, file_bytes: bytes, filename: str) -> str:
        from io import BytesIO
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
                text_parts = []
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    rows = list(sheet.iter_rows(values_only=True))
                    if not rows:
                        continue
                    text_parts.append(f"Sheet: {sheet_name}")
                    for row in rows:
                        if not row or all(c is None or str(c).strip() == "" for c in row):
                            continue
                        row_str = " | ".join(str(c).strip() for c in row if c is not None and str(c).strip() != "")
                        if row_str:
                            text_parts.append(row_str)
                wb.close()
                return "\n".join(text_parts)
            except Exception as exc:
                logger.warning(f"openpyxl failed for '{filename}', trying pandas fallback: {exc}")

        try:
            import pandas as pd
            excel_file = pd.ExcelFile(BytesIO(file_bytes))
            text_parts = []
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                if df.empty:
                    continue
                text_parts.append(f"Sheet: {sheet_name}")
                text_parts.append(df.to_string(index=False))
            return "\n".join(text_parts)
        except Exception as exc:
            raise Exception(f"Excel parsing failed: {exc}")

    def _extract_csv(self, file_bytes: bytes) -> str:
        import csv
        import io

        content = file_bytes.decode("utf-8", errors="ignore")
        reader = csv.reader(io.StringIO(content))
        lines = []
        for row in reader:
            if any(cell.strip() for cell in row):
                lines.append(" | ".join(cell.strip() for cell in row if cell.strip()))
        return "\n".join(lines)

    def chunk_text(self, text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
        """Sliding-window chunker breaking cleanly on sentence/word boundaries."""
        normalized = " ".join(text.split())
        if not normalized:
            return []

        chunks: List[str] = []
        start = 0
        text_len = len(normalized)
        while start < text_len:
            end = min(start + chunk_size, text_len)
            if end < text_len:
                # Try breaking at a sentence boundary or word boundary
                boundary = max(normalized.rfind('. ', start, end), normalized.rfind(' ', start, end))
                if boundary > start + (chunk_size // 2):
                    end = boundary + 1

            chunk = normalized[start:end].strip()
            if chunk:
                chunks.append(chunk)

            start = end if end >= text_len else max(end - overlap, start + 1)
        return chunks

    def parse(self, session_id: str, filename: str, file_bytes: bytes) -> List[DocumentChunk]:
        """Full pipeline: validate -> extract -> chunk -> build DocumentChunk objects."""
        self.validate(filename, file_bytes)
        raw_text = self.extract_text(filename, file_bytes)

        if not raw_text.strip():
            raise DocumentParsingException(filename, "No extractable text found in document.")

        text_chunks = self.chunk_text(raw_text)
        logger.info(f"Parsed '{filename}' into {len(text_chunks)} chunks for session {session_id}.")

        chunks: List[DocumentChunk] = []
        for idx, chunk_text in enumerate(text_chunks):
            # Deterministic ID: same file + same chunk index re-uploaded -> same ID (idempotent upsert)
            raw_id = f"{session_id}:{filename}:{idx}"
            doc_id = hashlib.sha256(raw_id.encode("utf-8")).hexdigest()
            chunks.append(
                DocumentChunk(
                    doc_id=doc_id,
                    session_id=session_id,
                    source_filename=filename,
                    chunk_index=idx,
                    text=chunk_text,
                )
            )
        return chunks